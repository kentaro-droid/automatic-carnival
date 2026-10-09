"""Regression tests using synthetic workbooks; no customer data or display needed."""
from pathlib import Path
import re
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
from zipfile import ZIP_DEFLATED, ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from app import App, EDITABLE_FIELDS
from validator import validate_property
from workbook_io import MASTER_SHEET, load_master, save_updates_preserving_validations


class MissingIdTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / 'source.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = MASTER_SHEET
        for col, header in enumerate(['物件ID', '物件名_日本語', '賃料_円'], 1):
            ws.cell(4, col, header)
        ws.cell(5, 2, 'テスト物件')  # A5 is absent, not merely an empty existing cell.
        ws.cell(5, 3, 500000)
        ws.cell(6, 2).number_format = '@'  # Styled, completely empty row.
        ws.cell(7, 1, 'P001')
        ws.cell(8, 1, 'P001')
        ws.cell(9, 2, ' \t ')
        ws.cell(10, 3, 0)
        dv = DataValidation(type='list', formula1='"A,B"')
        ws.add_data_validation(dv)
        dv.add('B5:B20')
        wb.save(self.source)
        wb.close()
        self.rows, self.headers = load_master(str(self.source))

    def model(self):
        # Exercise the actual GUI handlers without constructing a Tk window.
        model = Mock()
        model.rows = self.rows
        model.pending = {}
        model.current = self.rows[0]
        model.merged_row = lambda row: App.merged_row(model, row)
        model.duplicate_ids = lambda: App.duplicate_ids(model)
        model.display = lambda value: App.display(model, value)
        model.parse = lambda field, value: App.parse(model, field, value)
        return model

    def test_missing_id_is_loaded_and_validated(self):
        self.assertEqual([r['excel_row'] for r in self.rows], [5, 7, 8, 10])
        self.assertIn('ID_MISSING', [i.code for i in validate_property(self.rows[0])])
        self.assertEqual(self.rows[-1]['賃料_円'], 0)

    def test_empty_and_whitespace_rows_do_not_generate_errors(self):
        self.assertNotIn(6, [r['excel_row'] for r in self.rows])
        self.assertNotIn(9, [r['excel_row'] for r in self.rows])

    def test_duplicate_ids_remain_errors_on_both_rows(self):
        model = self.model()
        self.assertEqual(model.duplicate_ids(), {'P001'})
        for row in self.rows[1:3]:
            self.assertIn('ID_DUPLICATE', [i.code for i in validate_property(row, model.duplicate_ids())])

    def test_gui_lists_missing_id_and_can_complete_it(self):
        self.assertIn('物件ID', EDITABLE_FIELDS)
        model = self.model()
        model.tree.get_children.return_value = []
        App.populate_tree(model)
        self.assertEqual(model.tree.insert.call_args_list[0].kwargs['iid'], '5')
        self.assertIn('未入力', model.tree.insert.call_args_list[0].kwargs['values'][0])
        model.vars = {'物件ID': Mock()}
        model.vars['物件ID'].get.return_value = ' P001 '
        with patch('app.messagebox.showinfo'):
            App.apply_form(model)
        self.assertIn('ID_DUPLICATE', [i.code for i in validate_property(model.merged_row(model.current), model.duplicate_ids())])
        model.vars['物件ID'].get.return_value = 'P002'
        with patch('app.messagebox.showinfo'):
            App.apply_form(model)
        codes = [i.code for i in validate_property(model.merged_row(model.current), model.duplicate_ids())]
        self.assertNotIn('ID_MISSING', codes)
        self.assertNotIn('ID_DUPLICATE', codes)
        self.assertEqual(model.pending[5]['物件ID'], 'P002')
        self.assertIsNone(self.rows[0]['物件ID'])

    def test_numeric_existing_id_conflicts_with_edited_text(self):
        model = self.model()
        model.rows = [{'excel_row': 5, '物件ID': 123}, {'excel_row': 7, '物件ID': ' 123 '}]
        self.assertEqual(model.duplicate_ids(), {'123'})
        for row in model.rows:
            self.assertIn('ID_DUPLICATE', [i.code for i in validate_property(row, model.duplicate_ids())])
        model.current = model.rows[0]
        model.vars = {'物件ID': Mock()}
        model.vars['物件ID'].get.return_value = '123'
        with patch('app.messagebox.showinfo'):
            App.apply_form(model)
        self.assertEqual(model.pending[5], {})  # No numeric-to-text rewrite on unchanged input.

    def test_id_save_preserves_standard_and_x14_validations(self):
        # Inject an x14 extension, which openpyxl would discard if used to save.
        ext = ('<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
               'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
               '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
               '<x14:dataValidation type="list"><x14:formula1><xm:f>"A,B"</xm:f></x14:formula1>'
               '<xm:sqref>B5:B20</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')
        with ZipFile(self.source) as z:
            entries = {name: z.read(name) for name in z.namelist()}
        sheet = 'xl/worksheets/sheet1.xml'
        entries[sheet] = entries[sheet].replace(b'</worksheet>', ext.encode() + b'</worksheet>')
        with ZipFile(self.source, 'w', ZIP_DEFLATED) as z:
            for name, content in entries.items():
                z.writestr(name, content)
        original = self.source.read_bytes()
        dest = Path(self.temp.name) / 'reviewed.xlsx'
        save_updates_preserving_validations(str(self.source), str(dest), [(5, 1, 'P002'), (7, 1, 'P003')])
        self.assertEqual(self.source.read_bytes(), original)
        with ZipFile(dest) as z:
            new = z.read(sheet).decode()
            self.assertIn(ext, new)
            old_dv = re.search(r'<dataValidations\b.*?</dataValidations>', entries[sheet].decode()).group()
            self.assertIn(old_dv, new)
            for name, content in entries.items():
                if name != sheet:
                    self.assertEqual(z.read(name), content)
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            wb = load_workbook(dest)
        self.assertEqual(wb[MASTER_SHEET]['A5'].value, 'P002')
        self.assertEqual(wb[MASTER_SHEET]['A7'].value, 'P003')
        wb.close()


if __name__ == '__main__':
    unittest.main()

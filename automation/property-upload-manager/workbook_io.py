from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import re
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

MASTER_SHEET = '01_物件マスター'
HEADER_ROW = 4
DATA_START_ROW = 5


def load_master(path: str):
    # IMPORTANT: read only. Saving with openpyxl would remove x14 list validations in this workbook.
    wb = load_workbook(path, data_only=False, read_only=False)
    ws = wb[MASTER_SHEET]
    headers = {ws.cell(HEADER_ROW, c).value: c for c in range(1, ws.max_column + 1)}
    rows = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        values = {h: ws.cell(r, c).value for h, c in headers.items() if h}
        # Ignore only empty records; zero and False are meaningful cell values.
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values.values()):
            continue
        rows.append({'excel_row': r, **values})
    wb.close()
    return rows, headers


def _sheet_xml_path(xlsx_path: str, sheet_name: str) -> str:
    ns_main = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ns_rel = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    ns_pkg = 'http://schemas.openxmlformats.org/package/2006/relationships'
    with ZipFile(xlsx_path) as z:
        root = ET.fromstring(z.read('xl/workbook.xml'))
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        relmap = {el.attrib['Id']: el.attrib['Target'] for el in rels.findall(f'{{{ns_pkg}}}Relationship')}
        for s in root.findall(f'.//{{{ns_main}}}sheet'):
            if s.attrib.get('name') == sheet_name:
                rid = s.attrib[f'{{{ns_rel}}}id']
                target = relmap[rid]
                return target.lstrip('/') if target.startswith('/') else 'xl/' + target
    raise KeyError(sheet_name)


def _cell_ref(row: int, col: int) -> str:
    from openpyxl.utils import get_column_letter
    return f'{get_column_letter(col)}{row}'


def _serialize_value(value) -> tuple[str | None, str]:
    """Return (cell type, inner XML) without changing worksheet namespace declarations."""
    if value is None or value == '':
        return None, ''
    if isinstance(value, (datetime, date)):
        return None, f'<v>{to_excel(value)}</v>'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return None, f'<v>{value}</v>'
    text = str(value)
    preserve = ' xml:space="preserve"' if text.startswith(' ') or text.endswith(' ') else ''
    return 'inlineStr', f'<is><t{preserve}>{escape(text)}</t></is>'


def _patch_existing_cell(xml: str, ref: str, value) -> tuple[str, bool]:
    """Patch one cell while preserving all workbook XML prefixes/extension markup verbatim."""
    cell_re = re.compile(
        rf'<c\b(?P<attrs_self>[^>]*\br="{re.escape(ref)}"[^>]*)/>|'
        rf'<c\b(?P<attrs>[^>]*\br="{re.escape(ref)}"[^>]*)>(?P<body>.*?)</c>',
        re.DOTALL,
    )
    m = cell_re.search(xml)
    if not m:
        return xml, False

    is_self = m.group('attrs') is None
    attrs = m.group('attrs') if not is_self else m.group('attrs_self')
    if is_self:
        attrs = attrs.rstrip().rstrip('/').rstrip()
    attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
    cell_type, inner = _serialize_value(value)
    if cell_type:
        attrs += f' t="{cell_type}"'
    replacement = f'<c{attrs}>{inner}</c>'
    return xml[:m.start()] + replacement + xml[m.end():], True


def _insert_cell_into_existing_row(xml: str, ref: str, value) -> str:
    row_num = int(''.join(ch for ch in ref if ch.isdigit()))
    row_re = re.compile(rf'(<row\b[^>]*\br="{row_num}"[^>]*>)(.*?)(</row>)', re.DOTALL)
    m = row_re.search(xml)
    if not m:
        raise ValueError(f'Excel row {row_num} was not found; refusing to rewrite worksheet structure.')
    cell_type, inner = _serialize_value(value)
    type_attr = f' t="{cell_type}"' if cell_type else ''
    new_cell = f'<c r="{ref}"{type_attr}>{inner}</c>'
    replacement = m.group(1) + m.group(2) + new_cell + m.group(3)
    return xml[:m.start()] + replacement + xml[m.end():]


def _patch_cell_value(xml: str, ref: str, value) -> str:
    xml2, found = _patch_existing_cell(xml, ref, value)
    if found:
        return xml2
    return _insert_cell_into_existing_row(xml, ref, value)


def save_updates_preserving_validations(source_path: str, dest_path: str, updates: list[tuple[int, int, object]]):
    """
    Patch only cell values in the raw worksheet XML.

    This deliberately avoids serializing the entire worksheet with ElementTree because that can rename
    namespace prefixes referenced by mc:Ignorable and damage Excel x14 data-validation extensions.
    """
    sheet_path = _sheet_xml_path(source_path, MASTER_SHEET)
    with ZipFile(source_path, 'r') as zin:
        sheet_xml = zin.read(sheet_path).decode('utf-8')
        for row, col, value in updates:
            sheet_xml = _patch_cell_value(sheet_xml, _cell_ref(row, col), value)
        new_sheet = sheet_xml.encode('utf-8')
        Path(dest_path).parent.mkdir(parents=True, exist_ok=True)
        with ZipFile(dest_path, 'w', ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = new_sheet if item.filename == sheet_path else zin.read(item.filename)
                zout.writestr(item, data)
